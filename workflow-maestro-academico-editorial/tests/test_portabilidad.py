"""Pruebas ejecutables en Linux, Windows y macOS, sin descargas."""
import json
import importlib.util
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[2]
PREPROCESS = ROOT / 'preprocesador-documentos/scripts/preprocesar_documento.py'
RUNNER = ROOT / 'workflow-maestro-academico-editorial/scripts/ejecutar.py'


class PortabilidadTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='skills pruebas á ')
        self.addCleanup(self.tmp.cleanup)
        self.work = Path(self.tmp.name)
        self.source = self.work / 'capítulo uno.md'
        self.source.write_text('# Investigación\n\nTexto con educación y análisis.\n\n'
                               '| Indicador | Valor |\n| --- | --- |\n| Muestra | 12 |\n', encoding='utf-8')

    def run_script(self, script, *args, **kwargs):
        env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONIOENCODING='ascii')
        return subprocess.run([sys.executable, '-X', 'utf8', str(script), *map(str, args)],
                              cwd=self.work, env=env, capture_output=True,
                              encoding='utf-8', timeout=30, **kwargs)

    def test_preprocess_text_without_dependencies_from_other_directory(self):
        output = self.work / 'salida con ñ'
        result = self.run_script(PREPROCESS, self.source, output)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(len(list(output.iterdir())), 6)
        data = json.loads((output / 'capítulo uno.json').read_text(encoding='utf-8'))
        self.assertIn('educación', data['markdown'])
        self.assertTrue((output / 'capítulo uno-protegidos.json').is_file())

    def test_default_output_next_to_input(self):
        result = self.run_script(PREPROCESS, self.source)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((self.work / 'capítulo uno-preprocesado/capítulo uno.md').is_file())

    def test_rejects_invalid_limit_without_creating_output(self):
        output = self.work / 'invalid'
        result = self.run_script(PREPROCESS, self.source, output, '--max-words', '99')
        self.assertEqual(result.returncode, 2)
        self.assertFalse(output.exists())

    def test_rejects_missing_input_without_creating_output(self):
        output = self.work / 'invalid'
        result = self.run_script(PREPROCESS, self.work / 'missing.pdf', output)
        self.assertEqual(result.returncode, 2)
        self.assertFalse(output.exists())

    def test_refuses_to_overwrite_input(self):
        original = self.source.read_bytes()
        result = self.run_script(PREPROCESS, self.source, self.work)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(self.source.read_bytes(), original)

    def test_runner_resolves_skill_script_from_its_own_directory(self):
        result = self.run_script(RUNNER, 'auditor-documental-academico/scripts/inventariar_documento.py', self.source)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('Investigación', result.stdout)

    def test_runner_preserves_failure_exit_code(self):
        result = self.run_script(RUNNER, 'auditor-documental-academico/scripts/inventariar_documento.py')
        self.assertEqual(result.returncode, 2, result.stderr)

    @unittest.skipUnless(importlib.util.find_spec('docx') and importlib.util.find_spec('fitz'),
                         'Requiere las dependencias PDF/DOCX')
    def test_docx_and_pdf_round_trip(self):
        from docx import Document
        import pymupdf as fitz
        docx = self.work / 'documento con ñ.docx'
        conversion = self.run_script(RUNNER, 'maquetacion-academica-preentrega/scripts/markdown_a_docx.py',
                                     self.source, '--out', docx)
        self.assertEqual(conversion.returncode, 0, conversion.stderr)
        document = Document(docx)
        self.assertEqual(len(document.tables), 1)
        self.assertIn('educación', '\n'.join(p.text for p in document.paragraphs))
        pdf = self.work / 'investigación.pdf'
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((72, 72), 'Educación y análisis documental.')
            document.save(pdf)
        for source in (docx, pdf):
            with self.subTest(source=source.name):
                output = self.work / (source.stem + '-resultado')
                result = self.run_script(PREPROCESS, source, output)
                self.assertEqual(result.returncode, 0, result.stderr)
                data = json.loads((output / (source.stem + '.json')).read_text(encoding='utf-8'))
                self.assertIn('educación', data['markdown'].lower())

    def test_native_wrapper(self):
        import shutil
        if os.name == 'nt':
            shell = shutil.which('powershell') or shutil.which('pwsh')
            if not shell:
                self.skipTest('PowerShell no está instalado')
            command = [shell, '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File',
                       str(PREPROCESS.with_suffix('.ps1')), str(self.source)]
        else:
            command = ['sh', str(PREPROCESS.with_suffix('.sh')), str(self.source)]
        result = subprocess.run(command, cwd=self.work, capture_output=True, encoding='utf-8', timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((self.work / 'capítulo uno-preprocesado/capítulo uno.json').exists())


if __name__ == '__main__':
    unittest.main()
