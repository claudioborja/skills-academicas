#!/usr/bin/env python3
"""Comprobaciones parciales APA 7; no certifica cumplimiento integral."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document._body)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document._body)


def inherited(paragraph, attribute):
    """Resolver formato directo y cadena de estilos, sin confundir cero con None."""
    value = getattr(paragraph.paragraph_format, attribute)
    style = paragraph.style
    while value is None and style is not None:
        value = getattr(style.paragraph_format, attribute)
        style = style.base_style
    return value


def audit(path: Path, markdown: Path | None, svg_dir: Path | None,
          language: str | None, margin_cm: float, line_spacing: float,
          min_analysis_words: int) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = [
        "Revisión parcial: comprobar manualmente portada, resumen, encabezado/paginación, "
        "citas y referencias, orden alfabético, apéndices, notas y renderizado. "
        "Cero errores mecánicos no certifica APA 7."
    ]
    document = Document(path)
    normal = document.styles["Normal"]
    text = "\n".join(p.text for p in document.paragraphs)

    if language and document.core_properties.language != language:
        errors.append(f"Perfil editorial: metadatos distintos de {language}.")
    if not text.strip() and not document.tables and not document.inline_shapes:
        errors.append("Documento vacío.")
    if any(mark in text for mark in ("�", "Ã", "Â", "â€")):
        errors.append("Texto con codificación dañada.")
    for section in document.sections:
        for name, value in (("superior", section.top_margin), ("inferior", section.bottom_margin),
                            ("izquierdo", section.left_margin), ("derecho", section.right_margin)):
            if abs(value.cm - margin_cm) > 0.03:
                errors.append(f"Margen {name} distinto de {margin_cm:.2f} cm.")

    recommended = {"Times New Roman": 12, "Calibri": 11, "Arial": 11,
                   "Georgia": 11, "Lucida Sans Unicode": 10, "Computer Modern": 10}
    name = normal.font.name
    size = normal.font.size.pt if normal.font.size else None
    if name not in recommended or size != recommended[name]:
        warnings.append("Verificar legibilidad y consistencia de la fuente base; no coincide "
                        "con los ejemplos habituales APA. No implica por sí solo incumplimiento.")

    in_references = False
    resource_title = False
    in_abstract = False
    for index, paragraph in enumerate(document.paragraphs, 1):
        value = paragraph.text.strip()
        if not value:
            continue
        style = paragraph.style.name if paragraph.style else ""
        heading = re.fullmatch(r"Heading ([1-5])", style)
        if style.startswith("Heading") or value.casefold() in {"referencias", "references", "bibliografía"}:
            in_references = value.casefold() in {"referencias", "references", "bibliografía"}
            in_abstract = value.casefold() in {"resumen", "abstract"}
        spacing = inherited(paragraph, "line_spacing")
        if spacing != line_spacing:
            errors.append(f"Interlineado distinto de {line_spacing} en párrafo {index}.")
        if style in {'APA Cover Title', 'APA Cover Data', 'APA Cover Note'}:
            # La portada tiene excepciones de alineación y espaciado respecto al cuerpo.
            # Sus datos y disposición siguen en la revisión manual declarada arriba.
            continue
        for attr in ("space_before", "space_after"):
            if inherited(paragraph, attr) not in (None, 0):
                errors.append(f"Espacio adicional ({attr}) en párrafo {index}; revisar excepciones.")
        if heading:
            level = int(heading.group(1))
            expected_alignment = 1 if level == 1 else 0
            if inherited(paragraph, "alignment") != expected_alignment:
                errors.append(f"Heading {level} no tiene alineación APA en párrafo {index}.")
            runs = [run for run in paragraph.runs if run.text.strip()]
            heading_runs = runs[:1] if level >= 4 else runs
            for run in heading_runs:
                bold = run.bold if run.bold is not None else paragraph.style.font.bold
                italic = run.italic if run.italic is not None else paragraph.style.font.italic
                if not bold or bool(italic) != (level in {3, 5}):
                    errors.append(f"Heading {level} no tiene énfasis APA en párrafo {index}.")
                    break
            indent = inherited(paragraph, "first_line_indent") or 0
            if abs(indent - (457200 if level >= 4 else 0)) > 10800:
                errors.append(f"Sangría incorrecta en Heading {level}, párrafo {index}.")
            if level >= 4 and (len(runs) < 2 or not runs[0].text.rstrip().endswith(".")):
                warnings.append(f"Revisar título integrado y punto final del nivel {level}, párrafo {index}.")
            resource_title = False
            continue
        if value.casefold() in {"referencias", "references", "bibliografía"}:
            continue
        label = bool(re.fullmatch(r"(?:Tabla|Figura|Table|Figure) [A-Z]?\d+", value))
        special = (style.startswith("Heading") or style.startswith("List")
                   or style in {"Title", "Subtitle", "Quote", "Intense Quote"}
                   or label or resource_title or in_abstract
                   or value.startswith(("Nota.", "Note.", "Palabras clave:", "Keywords:")))
        resource_title = label
        if in_references:
            left = inherited(paragraph, "left_indent")
            first = inherited(paragraph, "first_line_indent")
            if left is None or first is None or abs(left.cm - 1.27) > .03 or abs(first.cm + 1.27) > .03:
                errors.append(f"Referencia sin sangría francesa en párrafo {index}.")
        elif not special:
            indent = inherited(paragraph, "first_line_indent")
            if indent is None or abs(indent.cm - 1.27) > .03:
                errors.append(f"Sangría inicial incorrecta en párrafo {index}.")
            if inherited(paragraph, "alignment") not in (None, 0):
                errors.append(f"Alineación del cuerpo distinta de izquierda en párrafo {index}.")

    for number, table in enumerate(document.tables, 1):
        borders = table._tbl.tblPr.find(qn("w:tblBorders"))
        if borders is None:
            warnings.append(f"Tabla {number}: revisar bordes heredados y legibilidad visualmente.")
            continue
        for edge in ("left", "right", "insideV"):
            node = borders.find(qn(f"w:{edge}"))
            if node is not None and node.get(qn("w:val")) not in {"nil", "none"}:
                errors.append(f"Tabla {number} conserva borde {edge}.")
        for row_number, row in enumerate(table.rows, 1):
            props = row._tr.trPr
            if props is None or props.find(qn("w:cantSplit")) is None:
                warnings.append(f"Tabla {number}, fila {row_number}: revisar partición entre páginas.")

    sequence = list(blocks(document))
    for index, block in enumerate(sequence):
        if not isinstance(block, Paragraph):
            continue
        label = block.text.strip()
        match = re.fullmatch(r"(Tabla|Figura|Table|Figure) ([A-Z]?\d+)", label)
        if not match:
            continue
        if not block.runs or not all(run.bold for run in block.runs if run.text.strip()):
            errors.append(f"{label} no está en negrita.")
        if index + 2 >= len(sequence):
            errors.append(f"{label} tiene secuencia incompleta.")
            continue
        title, resource = sequence[index + 1:index + 3]
        if not isinstance(title, Paragraph) or not title.text.strip() or not all(run.italic for run in title.runs if run.text.strip()):
            errors.append(f"Título de {label} no está en cursiva.")
        if match.group(1) in {"Tabla", "Table"} and not isinstance(resource, Table):
            errors.append(f"{label} no precede tabla nativa.")
        if match.group(1) in {"Figura", "Figure"} and (not isinstance(resource, Paragraph) or not resource._p.xpath(".//w:drawing")):
            errors.append(f"{label} no precede figura incrustada.")
        if min_analysis_words > 0:
            cursor = index + 3
            if cursor < len(sequence) and isinstance(sequence[cursor], Paragraph) and sequence[cursor].text.startswith(("Nota.", "Note.")):
                cursor += 1
            analysis = sequence[cursor] if cursor < len(sequence) else None
            words = len(re.findall(r"\b\w+\b", analysis.text)) if isinstance(analysis, Paragraph) else 0
            if words < min_analysis_words or (isinstance(analysis, Paragraph) and analysis.style.name.startswith(("Heading", "List"))):
                errors.append(f"Perfil editorial: {label} tiene análisis posterior insuficiente ({words} palabras).")

    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        if language and f'w:val="{language}"' not in xml:
            warnings.append(f"Perfil editorial: verificar lengua heredada de las corridas ({language}).")
        embedded_svg = [n for n in archive.namelist() if n.startswith("word/media/") and n.lower().endswith(".svg")]
        embedded_png = [n for n in archive.namelist() if n.startswith("word/media/") and n.lower().endswith(".png")]
        svg_blips = xml.count("asvg:svgBlip")
        if embedded_svg and (len(embedded_png) < len(embedded_svg) or svg_blips != len(embedded_svg)):
            errors.append(
                f"SVG incrustados sin respaldo OOXML válido: {len(embedded_svg)} SVG, "
                f"{len(embedded_png)} PNG y {svg_blips} referencias svgBlip."
            )
        if markdown:
            source = markdown.read_text(encoding="utf-8")
            requested = re.findall(r"!\[[^]]*\]\(([^)]+\.svg)\)", source, re.I)
            if len(requested) != len(embedded_svg):
                errors.append(f"SVG solicitados: {len(requested)}; SVG incrustados: {len(embedded_svg)}.")

    if svg_dir:
        for svg in svg_dir.glob("*.svg"):
            source = svg.read_text(encoding="utf-8", errors="replace")
            if re.search(r"<text\b[^>]*class=[\"']title[\"']", source, re.I):
                errors.append(f"{svg.name} contiene título editorial visible.")
            if re.search(r"marker-end\s*(?:=|:)", source, re.I):
                errors.append(f"{svg.name} usa marker-end; sustituir por puntas explícitas para evitar fallos de renderizado.")
            if re.search(r"<path\b[^>]*class=[\"'](?:a|arrow)[\"'][^>]*>", source, re.I):
                errors.append(f"{svg.name} deja el trazo de conectores solo en CSS; usar fill/stroke explícitos.")

    return errors, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--markdown", type=Path)
    parser.add_argument("--svg-dir", type=Path)
    parser.add_argument("--language", help="Control editorial opcional, no requisito APA")
    parser.add_argument("--margin-cm", default=2.54, type=float)
    parser.add_argument("--line-spacing", default=2.0, type=float)
    parser.add_argument("--min-analysis-words", default=0, type=int,
                        help="Cuota editorial opcional; APA no prescribe un mínimo")
    args = parser.parse_args()
    errors, warnings = audit(args.docx, args.markdown, args.svg_dir, args.language,
                              args.margin_cm, args.line_spacing, args.min_analysis_words)
    for item in errors:
        print(f"ERROR: {item}")
    for item in warnings:
        print(f"ADVERTENCIA: {item}")
    print(f"Resultado parcial: {len(errors)} errores, {len(warnings)} advertencias; requiere revisión manual")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
