#!/usr/bin/env python3
"""Convertir Markdown editorial a DOCX con tablas nativas editables."""

from __future__ import annotations

import argparse
import binascii
import io
import json
import re
import struct
import sys
import zipfile
import zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'editor-en-jefe/scripts'))
from archivos_seguros import atomic_write, validate_outputs
from portada_apa import add_cover, validate_metadata

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor
except ImportError as exc:
    raise SystemExit("Falta python-docx. Ejecuta el lanzador de la colección con --preparar o instala python-docx en tu entorno.") from exc


TABLE_SEPARATOR_RE = re.compile(r"^\s*\|?(?:\s*:?-{3,}:?\s*\|)+\s*$")
IMAGE_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
LINK_RE = re.compile(r"\[([^]]+)\]\(([^)]+)\)")
INLINE_EMPHASIS_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
DOCUMENT_LANGUAGE = "es-EC"


def clean_inline(value: str) -> str:
    value = LINK_RE.sub(lambda match: f"{match.group(1)} ({match.group(2)})", value)
    return re.sub(r"(\*\*|\*|`)", "", value).strip()


def add_inline_runs(paragraph, value: str) -> None:
    """Añadir texto conservando cursivas y negritas Markdown sin dañar URL."""
    value = LINK_RE.sub(lambda match: f"{match.group(1)} ({match.group(2)})", value)
    cursor = 0
    for match in INLINE_EMPHASIS_RE.finditer(value):
        if match.start() > cursor:
            paragraph.add_run(value[cursor:match.start()].replace("`", ""))
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(value):
        paragraph.add_run(value[cursor:].replace("`", ""))


def png_chunk(kind: bytes, data: bytes) -> bytes:
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", binascii.crc32(kind + data) & 0xFFFFFFFF)


def svg_placeholder_png(seed: int) -> bytes:
    """Crear un PNG mínimo solo como último recurso de compatibilidad."""
    rgba = bytes(((seed * 53) % 255, (seed * 97) % 255, (seed * 193) % 255, 255))
    header = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + png_chunk(b"IHDR", header)
        + png_chunk(b"IDAT", zlib.compress(b"\x00" + rgba))
        + png_chunk(b"IEND", b"")
    )


def svg_fallback_png(path: Path, seed: int) -> bytes:
    """Renderizar una vista PNG real; Word la usa si no admite SVG nativo."""
    try:
        import pymupdf as fitz

        source = fitz.open(stream=path.read_bytes(), filetype="svg")
        page = source[0]
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False)
        return pixmap.tobytes("png")
    except Exception:
        return svg_placeholder_png(seed)


def svg_aspect_ratio(path: Path) -> float:
    head = path.read_text(encoding="utf-8", errors="replace")[:1000]
    viewbox = re.search(r"viewBox=[\"']\s*[-\d.]+\s+[-\d.]+\s+([\d.]+)\s+([\d.]+)", head, re.I)
    if viewbox and float(viewbox.group(1)):
        return float(viewbox.group(2)) / float(viewbox.group(1))
    width = re.search(r"\bwidth=[\"']([\d.]+)", head, re.I)
    height = re.search(r"\bheight=[\"']([\d.]+)", head, re.I)
    if width and height and float(width.group(1)):
        return float(height.group(1)) / float(width.group(1))
    return 0.6


def embed_svg_parts(document_bytes: bytes, mappings: list[tuple[str, Path, str]]) -> bytes:
    """Añadir SVG nativo conservando un PNG de respaldo compatible con Word."""
    with zipfile.ZipFile(io.BytesIO(document_bytes), "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}

    relationships_name = "word/_rels/document.xml.rels"
    relationships = entries[relationships_name].decode("utf-8")
    content_types = entries["[Content_Types].xml"].decode("utf-8")
    document_xml = entries["word/document.xml"].decode("utf-8")
    replacements: dict[str, bytes] = {}
    used_ids = [int(value) for value in re.findall(r'Id="rId(\d+)"', relationships)]
    next_relationship_id = max(used_ids, default=0) + 1

    for index, (placeholder, svg_path, alt_text) in enumerate(mappings, start=1):
        svg_name = f"figure-{index}.svg"
        svg_part = f"word/media/{svg_name}"
        fallback_match = re.search(
            rf'<Relationship Id="(rId\d+)" Type="[^"]+/image" Target="media/{re.escape(placeholder)}"/?>',
            relationships,
        )
        if fallback_match is None:
            raise RuntimeError(f"No se encontró la relación PNG para {placeholder}.")
        fallback_id = fallback_match.group(1)
        svg_id = f"rId{next_relationship_id}"
        next_relationship_id += 1
        svg_relationship = (
            f'<Relationship Id="{svg_id}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            f'Target="media/{svg_name}"/>'
        )
        relationships = relationships.replace("</Relationships>", svg_relationship + "</Relationships>")
        original_blip = f'<a:blip r:embed="{fallback_id}"/>'
        svg_blip = (
            f'<a:blip r:embed="{fallback_id}"><a:extLst>'
            '<a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}">'
            '<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
            f'r:embed="{svg_id}"/></a:ext></a:extLst></a:blip>'
        )
        if original_blip not in document_xml:
            raise RuntimeError(f"No se encontró el dibujo de respaldo {fallback_id}.")
        document_xml = document_xml.replace(original_blip, svg_blip, 1)
        override = f'<Override PartName="/{svg_part}" ContentType="image/svg+xml"/>'
        if override not in content_types:
            content_types = content_types.replace("</Types>", override + "</Types>")
        replacements[svg_part] = svg_path.read_bytes()

    entries[relationships_name] = relationships.encode("utf-8")
    entries["[Content_Types].xml"] = content_types.encode("utf-8")
    entries["word/document.xml"] = document_xml.encode("utf-8")
    entries.update(replacements)

    temporary = io.BytesIO()
    with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return temporary.getvalue()


def split_row(line: str) -> list[str]:
    return [clean_inline(cell.strip()) for cell in line.strip().strip("|").split("|")]


def set_font_properties(font, font_name: str, font_size: float, language_code: str = DOCUMENT_LANGUAGE) -> None:
    """Fijar tipografía explícita, sin colores ni fuentes heredados del tema de Word."""
    font.name = font_name
    font.size = Pt(font_size)
    font.color.rgb = RGBColor(0, 0, 0)
    rpr = font._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attribute}"), font_name)
    language = rpr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        rpr.append(language)
    language.set(qn("w:val"), language_code)
    language.set(qn("w:eastAsia"), language_code)


def normalize_run(run, font_name: str, font_size: float, language: str = DOCUMENT_LANGUAGE) -> None:
    set_font_properties(run.font, font_name, font_size, language)


def repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    props.append(marker)


def prevent_row_split(row) -> None:
    props = row._tr.get_or_add_trPr()
    if props.find(qn("w:cantSplit")) is None:
        props.append(OxmlElement("w:cantSplit"))


def set_border(parent, edge: str, value: str = "single", size: str = "8") -> None:
    properties = parent._tc.get_or_add_tcPr() if hasattr(parent, "_tc") else parent._tbl.tblPr
    borders_tag = "w:tcBorders" if hasattr(parent, "_tc") else "w:tblBorders"
    borders = properties.find(qn(borders_tag))
    if borders is None:
        borders = OxmlElement(borders_tag)
        properties.append(borders)
    border = borders.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        borders.append(border)
    border.set(qn("w:val"), value)
    if value != "nil":
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")


def add_native_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    for edge in ("left", "right", "insideV", "insideH"):
        set_border(table, edge, "nil")
    for edge in ("top", "bottom"):
        set_border(table, edge)
    repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        prevent_row_split(table.rows[row_index])
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            cell.text = values[column_index] if column_index < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            cell.paragraphs[0].paragraph_format.first_line_indent = Inches(0)
            if row_index == 0:
                set_border(cell, "bottom")
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def configure_document(
    document, font_name: str, font_size: float, margin_inches: float,
    line_spacing: float, page_size: str, apa7_strict: bool = False,
) -> None:
    for section in document.sections:
        if page_size == "A4":
            section.page_width = Mm(210)
            section.page_height = Mm(297)
        elif page_size == "Letter":
            section.page_width = Inches(8.5)
            section.page_height = Inches(11)
        section.top_margin = Inches(margin_inches)
        section.bottom_margin = Inches(margin_inches)
        section.left_margin = Inches(margin_inches)
        section.right_margin = Inches(margin_inches)
    normal = document.styles["Normal"]
    set_font_properties(normal.font, font_name, font_size)
    normal.paragraph_format.line_spacing = line_spacing
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)
    normal.paragraph_format.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT if apa7_strict else WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    for style_name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3",
                       "Heading 4", "Heading 5", "Heading 6", "Heading 7",
                       "Heading 8", "Heading 9"):
        if style_name in document.styles:
            style = document.styles[style_name]
            set_font_properties(style.font, font_name, font_size)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.first_line_indent = Inches(0)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if apa7_strict:
                style.paragraph_format.line_spacing = 2
                style.font.italic = False
    if apa7_strict:
        document.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.styles["Title"].font.bold = True
        document.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.styles["Heading 1"].font.bold = True
        document.styles["Heading 2"].font.bold = True
        document.styles["Heading 3"].font.bold = True
        document.styles["Heading 3"].font.italic = True
        for level in (4, 5):
            style = document.styles[f"Heading {level}"]
            style.font.bold = True
            style.font.italic = level == 5
            style.paragraph_format.first_line_indent = Inches(.5)
        quote = document.styles["Quote"]
        set_font_properties(quote.font, font_name, font_size)
        quote.font.italic = False
        quote.font.bold = False
        quote.paragraph_format.line_spacing = 2
        quote.paragraph_format.space_before = Pt(0)
        quote.paragraph_format.space_after = Pt(0)
        for section in document.sections:
            # No destruir encabezados de plantillas. Completar solo los vacíos.
            for header in (section.header, section.first_page_header if section.different_first_page_header_footer else section.header,
                           section.even_page_header if document.settings.odd_and_even_pages_header_footer else section.header):
                if any(p.text.strip() or p._p.xpath(".//w:instrText|.//w:fldSimple") for p in header.paragraphs):
                    continue
                paragraph = header.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                paragraph.paragraph_format.first_line_indent = Inches(0)
                field = OxmlElement("w:fldSimple")
                field.set(qn("w:instr"), "PAGE")
                paragraph._p.append(field)


def resolve_image(source: Path, reference: str, assets_root: Path | None) -> Path:
    candidate = Path(reference)
    if candidate.is_absolute():
        return candidate
    local = (source.parent / candidate).resolve()
    if local.is_file() or assets_root is None:
        return local
    rooted = (assets_root / candidate).resolve()
    if rooted.is_file():
        return rooted
    # Al integrar capítulos en otra carpeta, una ruta relativa válida en el
    # capítulo puede dejar de serlo. El nombre del activo sigue siendo estable.
    return (assets_root / candidate.name).resolve()


def convert(
    source: Path,
    destination: Path,
    template: Path | None = None,
    assets_root: Path | None = None,
    font_name: str = "Times New Roman",
    font_size: float = 12,
    margin: float = 2.5 / 2.54,
    line_spacing: float = 1.5,
    page_size: str | None = None,
    image_width: float = 6.2,
    apa7_strict: bool = False,
    language: str = DOCUMENT_LANGUAGE,
    overwrite: bool = False,
    apa_metadata: dict | None = None,
) -> tuple[int, int]:
    validate_outputs([source, template], [destination], overwrite)
    if apa_metadata is not None:
        if not apa7_strict or template is not None:
            raise ValueError('La portada automática requiere --apa7-strict y no admite --template')
        validate_metadata(apa_metadata)
    lines = source.read_text(encoding="utf-8").splitlines()
    images = [resolve_image(source, match.group(2), assets_root)
              for line in lines if (match := IMAGE_RE.match(line.strip()))]
    validate_outputs([source, template, *images], [destination], overwrite)
    if apa7_strict:
        margin, line_spacing = 1.0, 2.0
    page_size = page_size or ("Letter" if apa7_strict else "A4")
    document = Document(str(template)) if template else Document()
    document.core_properties.language = language
    configure_document(document, font_name, font_size, margin, line_spacing, page_size, apa7_strict)
    if apa_metadata is not None:
        add_cover(document, apa_metadata, font_name, font_size, language)
    table_count = 0
    image_count = 0
    svg_mappings: list[tuple[str, Path, str]] = []
    index = 0
    in_references = False
    in_abstract = False
    expect_resource_title = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.casefold() in ("<!-- pagebreak -->", "\\pagebreak"):
            document.add_page_break()
            index += 1
            continue
        if index + 1 < len(lines) and "|" in line and TABLE_SEPARATOR_RE.match(lines[index + 1]):
            table_lines = [line]
            cursor = index + 2
            while cursor < len(lines) and "|" in lines[cursor] and lines[cursor].strip():
                table_lines.append(lines[cursor])
                cursor += 1
            add_native_table(document, [split_row(row) for row in table_lines])
            table_count += 1
            index = cursor
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            heading_text = clean_inline(heading.group(2))
            level = len(heading.group(1))
            if apa7_strict and level > 5:
                raise ValueError("APA tiene cinco niveles; reorganizar el encabezado de nivel 6.")
            paragraph = document.add_heading(heading_text, level=level)
            if apa7_strict and level in (4, 5):
                paragraph.clear()
                paragraph.paragraph_format.first_line_indent = Inches(.5)
                run = paragraph.add_run(heading_text.rstrip(".") + ".")
                run.bold, run.italic = True, level == 5
                cursor = index + 1
                while cursor < len(lines) and not lines[cursor].strip():
                    cursor += 1
                if cursor >= len(lines) or re.match(r"^(?:#|>|[-*+]\s|\d+[.)]\s|!\[|\||<!--|\\\\pagebreak)", lines[cursor].strip()):
                    raise ValueError("Los niveles APA 4 y 5 necesitan un párrafo de texto inmediatamente después.")
                paragraph.add_run(" ").bold = False
                start = len(paragraph.runs)
                add_inline_runs(paragraph, lines[cursor].strip())
                for run in paragraph.runs[start:]:
                    if run.bold is None:
                        run.bold = False
                    if run.italic is None:
                        run.italic = False
                index = cursor
            follows_abstract = in_abstract
            in_references = heading_text.casefold() in {"referencias", "bibliografía", "references"}
            in_abstract = heading_text.casefold() in {"resumen", "abstract"}
            if apa7_strict and (follows_abstract or in_references or in_abstract or re.match(r"^(Apéndice|Appendix)\b", heading_text, re.I)):
                paragraph.paragraph_format.page_break_before = True
            expect_resource_title = False
            index += 1
            continue
        resource_number = re.match(
            r"^(?:\*\*)?(Tabla|Figura|Gráfico)\s+([A-Za-z0-9.-]+)\.?(?:\*\*)?$",
            stripped,
            re.I,
        )
        if resource_number:
            paragraph = document.add_paragraph(f"{resource_number.group(1).title()} {resource_number.group(2)}")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.runs[0].bold = True
            expect_resource_title = True
            index += 1
            continue
        if expect_resource_title and re.match(r"^\*[^*].*\*$", stripped):
            paragraph = document.add_paragraph(clean_inline(stripped))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.italic = True
            expect_resource_title = False
            index += 1
            continue
        image = IMAGE_RE.match(stripped)
        if image:
            image_path = resolve_image(source, image.group(2), assets_root)
            if image_path.is_file():
                if image_path.suffix.casefold() == ".svg":
                    placeholder = io.BytesIO(svg_fallback_png(image_path, len(svg_mappings) + 1))
                    document.add_picture(
                        placeholder,
                        width=Inches(image_width),
                        height=Inches(image_width * svg_aspect_ratio(image_path)),
                    )
                    inline = document.inline_shapes[-1]._inline
                    blip = inline.xpath(".//a:blip")[0]
                    relationship_id = blip.get(qn("r:embed"))
                    image_part = document.part.related_parts[relationship_id]
                    placeholder_name = Path(str(image_part.partname)).name
                    svg_mappings.append((placeholder_name, image_path, clean_inline(image.group(1))))
                else:
                    document.add_picture(str(image_path), width=Inches(image_width))
                document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                document.paragraphs[-1].paragraph_format.first_line_indent = Inches(0)
                document.paragraphs[-1].paragraph_format.keep_with_next = True
                if image.group(1):
                    document.inline_shapes[-1]._inline.docPr.set("descr", clean_inline(image.group(1)))
                image_count += 1
            else:
                document.add_paragraph(f"[Imagen no encontrada: {image.group(2)}]")
            index += 1
            continue
        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        note = re.match(r"^\*Nota\.\*\s*(.*)$", stripped, re.I)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet.group(1))
        elif numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, numbered.group(1))
        elif note:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.space_before = Pt(0 if apa7_strict else 6)
            label = paragraph.add_run("Nota.")
            label.italic = True
            if note.group(1):
                paragraph.add_run(" ")
                add_inline_runs(paragraph, note.group(1))
        elif stripped.startswith(">"):
            paragraph = document.add_paragraph(style="Quote" if apa7_strict else "Intense Quote")
            parts = [stripped[1:].lstrip()]
            if apa7_strict:
                while index + 1 < len(lines) and lines[index + 1].strip().startswith(">") and lines[index + 1].strip() != ">":
                    index += 1
                    parts.append(lines[index].strip()[1:].lstrip())
                paragraph.paragraph_format.left_indent = Inches(.5)
                paragraph.paragraph_format.right_indent = Inches(0)
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(paragraph, " ".join(parts))
        elif stripped:
            paragraph = document.add_paragraph()
            add_inline_runs(paragraph, stripped)
            if in_references:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.left_indent = Inches(0.5)
                paragraph.paragraph_format.first_line_indent = Inches(-0.5)
            elif apa7_strict and in_abstract and not re.match(r"^(?:Palabras clave|Keywords):", clean_inline(stripped), re.I):
                paragraph.paragraph_format.first_line_indent = Inches(0)
        index += 1
    # Una plantilla puede introducir formato directo. Normalizar cada corrida al
    # final conserva negritas y cursivas, pero elimina fuentes y colores ajenos.
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            normalize_run(run, font_name, font_size, language)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        normalize_run(run, font_name, font_size, language)
    for node in document.styles.element.iter(qn("w:lang")):
        node.set(qn("w:val"), language)
        node.set(qn("w:eastAsia"), language)
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    if svg_mappings:
        data = embed_svg_parts(data, svg_mappings)
    atomic_write(destination, data, overwrite=overwrite)
    return table_count, image_count


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--template", type=Path, help="DOCX base con estilos, portada, encabezados y pies")
    parser.add_argument("--assets-root", type=Path, help="Directorio alternativo para resolver imágenes")
    parser.add_argument("--font", default="Times New Roman", help="Fuente base")
    parser.add_argument("--language", default=DOCUMENT_LANGUAGE, help="Idioma del documento, no requisito APA")
    parser.add_argument("--font-size", default=12, type=float, help="Tamaño de texto base en puntos")
    parser.add_argument("--margin", type=float, help="Márgenes en pulgadas; compatibilidad con usos anteriores")
    parser.add_argument("--margin-cm", default=2.5, type=float, help="Márgenes en centímetros")
    parser.add_argument("--line-spacing", default=1.5, type=float, help="Interlineado")
    parser.add_argument("--page-size", choices=("A4", "Letter", "template"),
                        help="Predeterminado: Letter en APA; A4 en perfil editorial")
    parser.add_argument("--image-width", default=6.2, type=float, help="Ancho de imagen en pulgadas")
    parser.add_argument("--apa7-strict", action="store_true",
                        help="Aplicar márgenes de 2,54 cm, doble espacio, alineación izquierda y jerarquía APA 7")
    parser.add_argument('--overwrite', action='store_true', help='Reemplazar una salida existente, nunca entradas ni plantilla')
    parser.add_argument('--apa-metadata', type=Path, help='JSON con perfil y datos reales de portada APA; incompatible con plantilla')
    args = parser.parse_args()
    try:
        validate_outputs([args.source, args.template, args.apa_metadata], [args.out], args.overwrite)
    except ValueError as exc:
        parser.error(str(exc))
    if not args.source.is_file():
        parser.error(f"No existe el archivo: {args.source}")
    if args.template and not args.template.is_file():
        parser.error(f"No existe la plantilla: {args.template}")
    metadata = None
    if args.apa_metadata:
        try:
            metadata = json.loads(args.apa_metadata.read_text(encoding='utf-8'))
            validate_metadata(metadata)
            if not args.apa7_strict or args.template:
                raise ValueError('La portada automática requiere --apa7-strict y no admite --template')
        except (OSError, ValueError) as exc:
            parser.error(str(exc))
    margin = args.margin if args.margin is not None else args.margin_cm / 2.54
    line_spacing = args.line_spacing
    if args.apa7_strict:
        margin = 1.0
        line_spacing = 2.0
    tables, images = convert(
        args.source.resolve(), args.out.absolute(),
        args.template.resolve() if args.template else None,
        args.assets_root.resolve() if args.assets_root else None,
        args.font, args.font_size,
        margin, line_spacing, args.page_size, args.image_width, args.apa7_strict, args.language,
        overwrite=args.overwrite,
        apa_metadata=metadata,
    )
    print(f"DOCX creado: {args.out}; tablas nativas: {tables}; imágenes: {images}")
    if args.apa7_strict:
        print("Formato base APA aplicado; revisar portada, resumen, encabezados, citas, referencias y renderizado. No es certificación integral.")


if __name__ == "__main__":
    main()
