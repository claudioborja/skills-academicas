"""Portada básica APA con una afiliación compartida y datos suministrados."""
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def validate_metadata(data):
    if not isinstance(data, dict) or data.get('profile') not in ('student', 'professional'):
        raise ValueError('profile debe ser student o professional')
    required = {'profile', 'title', 'authors', 'affiliation'}
    required |= {'course', 'instructor', 'due_date'} if data['profile'] == 'student' else {'running_head'}
    allowed = required | ({'author_note', 'author_note_label'} if data['profile'] == 'professional' else set())
    if set(data) - allowed:
        raise ValueError('Campos APA desconocidos o incompatibles: ' + ', '.join(sorted(set(data) - allowed)))
    for key in required:
        if not isinstance(data.get(key), str) or not data[key].strip() or any(c in data[key] for c in '\r\n\t'):
            raise ValueError(f'Falta un texto válido de una línea para {key}')
    if data['profile'] == 'professional':
        if len(data['running_head'].strip().upper()) > 50:
            raise ValueError('El encabezado abreviado excede 50 caracteres')
        notes = data.get('author_note', [])
        if not isinstance(notes, list) or any(not isinstance(p, str) or not p.strip() for p in notes):
            raise ValueError('author_note debe ser una lista de párrafos no vacíos')
        if 'author_note_label' in data and (not isinstance(data['author_note_label'], str) or not data['author_note_label'].strip()):
            raise ValueError('author_note_label debe ser texto no vacío')


def add_cover(document, data, font_name, font_size, language):
    """Solo usar en un documento nuevo, nunca sobre una plantilla editorial."""
    validate_metadata(data)
    for name, base in (('APA Cover Title', 'Title'), ('APA Cover Data', 'Normal'), ('APA Cover Note', 'Normal')):
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles[base]
        if name == 'APA Cover Title':
            # Title puede traer un borde decorativo en la plantilla de python-docx.
            # Anularlo aquí, sin modificar estilos de otros productos editoriales.
            borders = OxmlElement('w:pBdr')
            for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
                border = OxmlElement(f'w:{edge}')
                border.set(qn('w:val'), 'nil')
                borders.append(border)
            style.element.get_or_add_pPr().append(borders)
    def centered(text, bold=False):
        paragraph = document.add_paragraph(style='APA Cover Title' if bold else 'APA Cover Data')
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Inches(0)
        paragraph.paragraph_format.line_spacing = 2
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = False
        return paragraph

    title = centered(data['title'].strip(), True)
    title.paragraph_format.space_before = Pt(font_size * 6)
    title.paragraph_format.space_after = Pt(font_size * 2)
    for key in ('authors', 'affiliation'):
        centered(data[key].strip())
    if data['profile'] == 'student':
        for key in ('course', 'instructor', 'due_date'):
            centered(data[key].strip())
    elif data.get('author_note'):
        label = data.get('author_note_label', 'Nota de autor' if language.startswith('es') else 'Author Note')
        heading = centered(label, True)
        heading.paragraph_format.space_before = Pt(font_size * 2)
        for text in data['author_note']:
            paragraph = document.add_paragraph(text, style='APA Cover Note')
            paragraph.paragraph_format.keep_with_next = True
    document.paragraphs[-1].paragraph_format.keep_with_next = False
    document.add_page_break()

    section = document.sections[0]
    section.header_distance = Inches(.5)
    header = section.header.paragraphs[0]
    header.clear()
    header.paragraph_format.first_line_indent = Inches(0)
    header.paragraph_format.line_spacing = 1
    header.style.font.name = font_name
    header.style.font.size = Pt(font_size)
    if data['profile'] == 'professional':
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # En este documento nuevo, quitar el tabulador central heredado de Header.
        header.style.paragraph_format.tab_stops.clear_all()
        header.paragraph_format.tab_stops.add_tab_stop(section.page_width - section.left_margin - section.right_margin, WD_TAB_ALIGNMENT.RIGHT)
        header.add_run(data['running_head'].strip().upper() + '\t')
    else:
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    header._p.append(field)
