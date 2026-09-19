from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'maquetacion-academica-preentrega/scripts'))
import markdown_a_docx as exporter
from docx import Document


class ExportSafetyTests(unittest.TestCase):
    def test_direct_export_refuses_existing_output(self):
        for script, suffix in [('markdown_a_docx.py', '.docx'), ('markdown_a_txt_final.py', '.txt')]:
            with self.subTest(script=script), tempfile.TemporaryDirectory() as tmp:
                source, dest = Path(tmp)/'entrada.md', Path(tmp)/('salida'+suffix)
                source.write_text('# Título\nTexto.', encoding='utf-8')
                dest.write_bytes(b'Anterior')
                result = subprocess.run([sys.executable, str(ROOT/'maquetacion-academica-preentrega/scripts'/script), str(source), '--out', str(dest)], capture_output=True)
                self.assertEqual(dest.read_bytes(), b'Anterior')
                self.assertNotEqual(result.returncode, 0)

    def test_txt_default_never_overwrites_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp)/'entrada.txt'
            source.write_bytes(b'# Texto original')
            result = subprocess.run([sys.executable, str(ROOT/'maquetacion-academica-preentrega/scripts/markdown_a_txt_final.py'), str(source)], capture_output=True)
            self.assertEqual(source.read_bytes(), b'# Texto original')
            self.assertNotEqual(result.returncode, 0)

    def test_docx_template_is_never_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, template = Path(tmp)/'entrada.md', Path(tmp)/'plantilla.docx'
            source.write_text('Nuevo texto', encoding='utf-8')
            Document().save(template)
            original = template.read_bytes()
            try:
                exporter.convert(source, template, template=template)
            except ValueError:
                pass
            self.assertEqual(template.read_bytes(), original)

    def test_svg_failure_leaves_no_partial_document(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, dest, svg = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx', Path(tmp)/'imagen.svg'
            svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><rect width="10" height="10"/></svg>', encoding='utf-8')
            source.write_text('![Figura](imagen.svg)', encoding='utf-8')
            with patch.object(exporter, 'embed_svg_parts', side_effect=OSError('fallo de empaquetado')):
                with self.assertRaises(OSError):
                    exporter.convert(source, dest)
            self.assertFalse(dest.exists())

    def test_explicit_overwrite_generates_readable_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, dest = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx'
            source.write_text('Texto conservado', encoding='utf-8')
            dest.write_bytes(b'Anterior')
            result = subprocess.run([sys.executable, str(ROOT/'maquetacion-academica-preentrega/scripts/markdown_a_docx.py'), str(source), '--out', str(dest), '--overwrite'], capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn('Texto conservado', [p.text for p in Document(dest).paragraphs])

    def test_overwrite_never_replaces_docx_inputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, template, svg = Path(tmp)/'entrada.md', Path(tmp)/'plantilla.docx', Path(tmp)/'imagen.svg'
            source.write_text('![Figura](imagen.svg)', encoding='utf-8')
            svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>', encoding='utf-8')
            Document().save(template)
            for dest in (source, template, svg):
                with self.subTest(dest=dest):
                    original = dest.read_bytes()
                    with self.assertRaises(ValueError):
                        exporter.convert(source, dest, template=template, overwrite=True)
                    self.assertEqual(dest.read_bytes(), original)

    def test_svg_failure_preserves_authorized_existing_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, dest, svg = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx', Path(tmp)/'imagen.svg'
            svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>', encoding='utf-8')
            source.write_text('![Figura](imagen.svg)', encoding='utf-8')
            dest.write_bytes(b'Anterior')
            with patch.object(exporter, 'embed_svg_parts', side_effect=OSError('fallo de empaquetado')):
                with self.assertRaises(OSError):
                    exporter.convert(source, dest, overwrite=True)
            self.assertEqual(dest.read_bytes(), b'Anterior')
            self.assertEqual(set(Path(tmp).iterdir()), {source, dest, svg})

    def test_native_svg_is_published_without_named_temporary_collision(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, dest, svg = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx', Path(tmp)/'imagen.svg'
            old_temp = Path(tmp)/'salida.svg-tmp.docx'
            old_temp.write_bytes(b'Archivo ajeno')
            svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>', encoding='utf-8')
            source.write_text('![Figura](imagen.svg)', encoding='utf-8')
            exporter.convert(source, dest)
            self.assertEqual(old_temp.read_bytes(), b'Archivo ajeno')
            with zipfile.ZipFile(dest) as archive:
                self.assertIsNone(archive.testzip())
                self.assertEqual(archive.read('word/media/figure-1.svg'), svg.read_bytes())
            self.assertEqual(len(Document(dest).inline_shapes), 1)

    def test_txt_explicit_overwrite_and_source_protection(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, dest = Path(tmp)/'entrada.txt', Path(tmp)/'salida.txt'
            source.write_text('# Título\nTexto.', encoding='utf-8')
            original = source.read_bytes()
            dest.write_bytes(b'Anterior')
            command = [sys.executable, str(ROOT/'maquetacion-academica-preentrega/scripts/markdown_a_txt_final.py'), str(source), '--overwrite']
            result = subprocess.run([*command, '--out', str(dest)], capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dest.read_text(encoding='utf-8'), 'Título\n\nTexto.\n')
            result = subprocess.run(command, capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(source.read_bytes(), original)
