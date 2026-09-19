"""Regresiones con DOCX reales; no certifican revisión semántica ni visual."""
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from auditar_docx_apa7 import audit
from markdown_a_docx import convert


class ApaAuditTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="apa7 pruebas ")
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "ejemplo.docx"
        self.doc = Document()
        for section in self.doc.sections:
            section.top_margin = section.bottom_margin = Inches(1)
            section.left_margin = section.right_margin = Inches(1)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal.paragraph_format.line_spacing = 2
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.first_line_indent = Inches(.5)
        normal.paragraph_format.alignment = 0
        self.doc.add_paragraph("Un párrafo ordinario de prueba.")

    def result(self):
        self.doc.save(self.path)
        return audit(self.path, None, None, None, 2.54, 2.0, 0)

    def test_accepts_valid_fonts_without_imposing_language_or_unused_headings(self):
        for name, size in [("Calibri", 11), ("Arial", 11), ("Georgia", 11),
                           ("Lucida Sans Unicode", 10), ("Times New Roman", 12)]:
            with self.subTest(font=name):
                self.doc.styles["Normal"].font.name = name
                self.doc.styles["Normal"].font.size = Pt(size)
                errors, warnings = self.result()
                self.assertEqual(errors, [])
                self.assertTrue(warnings)  # siempre explicitar límites

    def test_checks_later_paragraph_spacing(self):
        self.doc.add_paragraph("Error posterior.").paragraph_format.line_spacing = 1
        self.assertTrue(any("interlineado" in e.lower() for e in self.result()[0]))

    def test_checks_body_indentation_and_alignment(self):
        p = self.doc.add_paragraph("Sin sangría, justificado.")
        p.paragraph_format.first_line_indent = Inches(0)
        p.alignment = 3
        errors = self.result()[0]
        self.assertTrue(any("sangría" in e.lower() for e in errors))
        self.assertTrue(any("alineación" in e.lower() for e in errors))

    def test_empty_document_is_reported_without_exception(self):
        p = self.doc.paragraphs[0]._element
        p.getparent().remove(p)
        self.assertTrue(any("vacío" in e.lower() for e in self.result()[0]))

    def test_table_does_not_require_note_or_word_quota(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.add_run("Tabla 1").bold = True
        p = self.doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.add_run("Resultados").italic = True
        self.doc.add_table(rows=2, cols=2)
        self.assertEqual(self.result()[0], [])

    def test_references_end_at_appendix(self):
        h = self.doc.add_heading("Referencias", 1)
        h.alignment = 1
        h.style.font.bold = True
        h.style.font.italic = False
        h.style.paragraph_format.first_line_indent = Inches(0)
        h.style.paragraph_format.space_before = h.style.paragraph_format.space_after = Pt(0)
        p = self.doc.add_paragraph("Autor, A. (2020). Libro. Editorial.")
        p.paragraph_format.left_indent = Inches(.5)
        p.paragraph_format.first_line_indent = Inches(-.5)
        self.doc.add_heading("Apéndice A", 1)
        self.doc.add_paragraph("Contenido del apéndice.")
        self.assertFalse(any("Referencia sin" in e for e in self.result()[0]))

    def test_invalid_used_level_five_is_detected(self):
        self.doc.add_heading("Título mal formado", 5)
        self.assertTrue(any("Heading 5" in e for e in self.result()[0]))

    def test_inherited_zero_spacing_is_not_replaced_by_normal(self):
        from docx.enum.style import WD_STYLE_TYPE
        child = self.doc.styles.add_style("Cuerpo heredado", WD_STYLE_TYPE.PARAGRAPH)
        child.base_style = self.doc.styles["Normal"]
        child.paragraph_format.line_spacing = 1
        self.doc.paragraphs[0].style = child
        self.assertTrue(any("interlineado" in e.lower() for e in self.result()[0]))

    def test_editorial_language_is_only_enforced_when_requested(self):
        self.doc.save(self.path)
        errors, _ = audit(self.path, None, None, "es-PE", 2.54, 2, 0)
        self.assertTrue(any("Perfil editorial" in e for e in errors))


class ApaConverterTests(unittest.TestCase):
    def make(self, text):
        tmp = tempfile.TemporaryDirectory(prefix="apa7 conversión ")
        self.addCleanup(tmp.cleanup)
        source = Path(tmp.name) / "fuente.md"
        target = Path(tmp.name) / "salida.docx"
        source.write_text(text, encoding="utf-8")
        convert(source, target, apa7_strict=True)
        return Document(target)

    def test_strict_api_applies_spacing_margins_letter_and_page_number(self):
        doc = self.make("Texto.")
        self.assertAlmostEqual(doc.sections[0].top_margin.inches, 1, places=2)
        self.assertAlmostEqual(doc.sections[0].page_width.inches, 8.5, places=2)
        self.assertEqual(doc.styles["Normal"].paragraph_format.line_spacing, 2)
        self.assertIn("PAGE", doc.sections[0].header._element.xml)

    def test_levels_four_and_five_are_run_in(self):
        doc = self.make("#### Subtema\n\nTexto cuatro.\n\n##### Detalle\n\nTexto cinco.")
        self.assertEqual([p.text for p in doc.paragraphs],
                         ["Subtema. Texto cuatro.", "Detalle. Texto cinco."])
        for index, p in enumerate(doc.paragraphs):
            self.assertTrue(p.runs[0].bold)
            self.assertEqual(bool(p.runs[0].italic), index == 1)
            self.assertFalse(bool(p.runs[-1].bold))
            self.assertAlmostEqual(p.paragraph_format.first_line_indent.inches, .5)

    def test_block_quote_is_not_decorative_and_joins_wrapped_lines(self):
        doc = self.make("> Primera línea\n> segunda línea (Autor, 2020, p. 2).")
        self.assertEqual(len(doc.paragraphs), 1)
        p = doc.paragraphs[0]
        self.assertEqual(p.text, "Primera línea segunda línea (Autor, 2020, p. 2).")
        self.assertAlmostEqual(p.paragraph_format.left_indent.inches, .5)
        self.assertEqual(p.paragraph_format.first_line_indent, 0)
        self.assertFalse(bool(p.style.font.italic))

    def test_references_start_new_page(self):
        doc = self.make("Texto.\n\n# Referencias\n\nAutor. (2020). *Libro*. Editorial.")
        self.assertTrue(doc.paragraphs[1].paragraph_format.page_break_before)

    def test_abstract_has_no_indent_and_body_retains_it(self):
        doc = self.make("# Resumen\n\nResumen breve.\n\n# Desarrollo\n\nTexto ordinario.")
        self.assertEqual(doc.paragraphs[1].paragraph_format.first_line_indent, 0)
        self.assertIsNone(doc.paragraphs[3].paragraph_format.first_line_indent)
        self.assertAlmostEqual(doc.styles["Normal"].paragraph_format.first_line_indent.inches, .5)

    def test_generated_five_levels_pass_mechanical_audit(self):
        doc = self.make("# Uno\n\nTexto.\n\n## Dos\n\nTexto.\n\n### Tres\n\nTexto.\n\n#### Cuatro\n\nTexto.\n\n##### Cinco\n\nTexto.")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "niveles.docx"
            doc.save(path)
            self.assertEqual(audit(path, None, None, None, 2.54, 2, 0)[0], [])

    def test_cli_font_language_and_scope_notice(self):
        import subprocess
        with tempfile.TemporaryDirectory(prefix="apa cli ") as tmp:
            source, target = Path(tmp) / "área.md", Path(tmp) / "salida.docx"
            source.write_text("Texto en español.", encoding="utf-8")
            script = Path(__file__).resolve().parents[1] / "scripts" / "markdown_a_docx.py"
            result = subprocess.run([sys.executable, str(script), str(source), "--out", str(target),
                                     "--apa7-strict", "--font", "Arial", "--font-size", "11",
                                     "--language", "es-PE"], capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(target)
            self.assertEqual(doc.core_properties.language, "es-PE")
            self.assertIn('w:val="es-PE"', doc.element.xml)
            self.assertEqual(doc.styles["Normal"].font.name, "Arial")
            self.assertIn("No es certificación integral", result.stdout)


if __name__ == "__main__":
    unittest.main()
