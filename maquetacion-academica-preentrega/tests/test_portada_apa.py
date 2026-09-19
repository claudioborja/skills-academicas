import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
from markdown_a_docx import convert, configure_document
from portada_apa import add_cover
from auditar_docx_apa7 import audit
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

STUDENT = {'profile': 'student', 'title': 'Aprendizaje universitario',
           'authors': 'Ana Pérez', 'affiliation': 'Departamento de Educación, Universidad de Prueba',
           'course': 'EDU 101: Aprendizaje', 'instructor': 'Luis López', 'due_date': '19 de septiembre de 2026'}


class CoverTests(unittest.TestCase):
    def test_cover_title_and_note_do_not_inherit_decorative_borders(self):
        # Reproduce a decorated dependency template independently of its version.
        doc = Document()
        props = doc.styles['Title'].element.get_or_add_pPr()
        for old in props.findall(qn('w:pBdr')):
            props.remove(old)
        borders = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:color'), '4F81BD')
        borders.append(bottom)
        props.append(borders)
        configure_document(doc, 'Times New Roman', 12, 1, 2, 'Letter', True)
        data = {k: STUDENT[k] for k in ('title', 'authors', 'affiliation')}
        data.update(profile='professional', running_head='Aprendizaje', author_note=['Nota de prueba.'])
        add_cover(doc, data, 'Times New Roman', 12, 'es-EC')
        for p in (doc.paragraphs[0], next(p for p in doc.paragraphs if p.text == 'Nota de autor')):
            layers = [p._p]
            style = p.style
            while style is not None:
                layers.append(style.element)
                style = style.base_style
            for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
                values = [node.get(qn('w:val')) for layer in layers
                          for node in layer.xpath(f'./w:pPr/w:pBdr/w:{edge}')]
                self.assertTrue(not values or values[0] in ('nil', 'none'), (p.text, edge, values))

    def test_professional_header_has_no_intermediate_inherited_tab(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, out = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx'
            source.write_text('Cuerpo.', encoding='utf-8')
            data = {k: STUDENT[k] for k in ('title', 'authors', 'affiliation')}
            data.update(profile='professional', running_head='Aprendizaje')
            convert(source, out, apa7_strict=True, apa_metadata=data)
            doc = Document(out)
            header = doc.sections[0].header.paragraphs[0]
            layers = [header._p]
            style = header.style
            while style is not None:
                layers.append(style.element)
                style = style.base_style
            tabs = {}
            for layer in reversed(layers):
                for tab in layer.xpath('./w:pPr/w:tabs/w:tab'):
                    tabs[int(tab.get(qn('w:pos')))] = tab.get(qn('w:val'))
            active = sorted((pos, alignment) for pos, alignment in tabs.items() if alignment != 'clear')
            self.assertEqual(active, [(9360, 'right')])

    def test_body_after_abstract_starts_new_page_in_apa_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp)/'entrada.md'
            for label in ('Resumen', 'Abstract'):
                for strict in (True, False):
                    with self.subTest(label=label, strict=strict):
                        source.write_text(f'# {label}\n\nSíntesis.\n\n*Keywords:* learning\n\n# Título del trabajo\n\nCuerpo.\n\n## Método\n\nProcedimiento.', encoding='utf-8')
                        out = Path(tmp)/f'{label}-{strict}.docx'
                        convert(source, out, apa7_strict=strict)
                        doc = Document(out)
                        body = next(p for p in doc.paragraphs if p.text == 'Título del trabajo')
                        method = next(p for p in doc.paragraphs if p.text == 'Método')
                        self.assertEqual(bool(body.paragraph_format.page_break_before), strict)
                        self.assertFalse(method.paragraph_format.page_break_before)
                        self.assertEqual([p.text for p in doc.paragraphs if p.text],
                                         [label, 'Síntesis.', 'Keywords: learning', 'Título del trabajo', 'Cuerpo.', 'Método', 'Procedimiento.'])

    def test_student_cover_contains_supplied_fields_and_page_break(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, out = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx'
            source.write_text('Texto del cuerpo.', encoding='utf-8')
            convert(source, out, apa7_strict=True, apa_metadata=STUDENT)
            doc = Document(out)
            nonempty = [p for p in doc.paragraphs if p.text.strip()]
            self.assertEqual([p.text for p in nonempty[:6]], [STUDENT[k] for k in ('title','authors','affiliation','course','instructor','due_date')])
            self.assertTrue(nonempty[0].runs[0].bold)
            self.assertTrue(all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in nonempty[:6]))
            self.assertTrue(doc.element.xpath('.//w:br[@w:type="page"]'))
            self.assertTrue(doc.sections[0].header._element.xpath('.//w:fldSimple[@w:instr="PAGE"]'))
            self.assertNotIn(STUDENT['title'], doc.sections[0].header.paragraphs[0].text)
            errors, warnings = audit(out, None, None, None, 2.54, 2, 0)
            self.assertEqual(errors, [])

    def test_professional_cover_has_running_head_and_author_note(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, out = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx'
            source.write_text('Cuerpo.', encoding='utf-8')
            data = {k: STUDENT[k] for k in ('title','authors','affiliation')}
            data.update(profile='professional', running_head='Aprendizaje', author_note=['Sin conflictos declarados por la autora.'])
            convert(source, out, apa7_strict=True, apa_metadata=data)
            doc = Document(out)
            self.assertIn('APRENDIZAJE', doc.sections[0].header.paragraphs[0].text)
            self.assertNotIn('Running head:', doc.sections[0].header.paragraphs[0].text)
            self.assertIn(data['author_note'][0], [p.text for p in doc.paragraphs])

    def test_missing_fields_and_long_header_fail_without_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, out = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx'
            source.write_text('Cuerpo.', encoding='utf-8')
            professional = {k: STUDENT[k] for k in ('title', 'authors', 'affiliation')}
            professional.update(profile='professional', running_head='X'*51)
            for data in ({'profile':'student'}, professional, dict(STUDENT, typo='dato')):
                with self.subTest(data=data), self.assertRaises(ValueError):
                    convert(source, out, apa7_strict=True, apa_metadata=data)
                self.assertFalse(out.exists())

    def test_template_and_non_apa_mode_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, out, template = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx', Path(tmp)/'base.docx'
            source.write_text('Cuerpo.', encoding='utf-8')
            Document().save(template)
            for options in ({'apa7_strict':False}, {'apa7_strict':True,'template':template}):
                with self.assertRaises(ValueError):
                    convert(source, out, apa_metadata=STUDENT, **options)
                self.assertFalse(out.exists())

    def test_cli_metadata_and_collision_protection(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, out, metadata = Path(tmp)/'entrada.md', Path(tmp)/'salida.docx', Path(tmp)/'datos.json'
            source.write_text('Cuerpo.', encoding='utf-8')
            metadata.write_text(json.dumps(STUDENT), encoding='utf-8')
            original = metadata.read_bytes()
            script = Path(__file__).resolve().parents[1]/'scripts/markdown_a_docx.py'
            cmd = [sys.executable, str(script), str(source), '--apa7-strict', '--apa-metadata', str(metadata)]
            result = subprocess.run([*cmd,'--out',str(out)], capture_output=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn(STUDENT['title'], [p.text for p in Document(out).paragraphs])
            result = subprocess.run([*cmd,'--out',str(metadata),'--overwrite'], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(metadata.read_bytes(), original)
