#!/usr/bin/env python3
from __future__ import annotations

import argparse
import html
import json
import re
import sys
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'editor-en-jefe/scripts'))
from archivos_seguros import atomic_write, validate_outputs


SPACE_RE = re.compile(r"[ \t]+")
BLANK_RE = re.compile(r"\n{3,}")


@dataclass
class Conversion:
    source: str
    source_type: str
    extracted_at: str
    chars: int
    words: int
    warnings: list[str]
    markdown: str


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = SPACE_RE.sub(" ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = BLANK_RE.sub("\n\n", text)
    return text.strip()


def read_pdf(path: Path) -> tuple[str, list[str]]:
    warnings: list[str] = []
    try:
        import pymupdf as fitz  # type: ignore

        doc = fitz.open(path)
        pages = []
        for idx, page in enumerate(doc, start=1):
            pages.append(f"\n\n<!-- page:{idx} -->\n\n{page.get_text('text')}")
        return "\n".join(pages), warnings
    except Exception as exc:
        warnings.append(f"PyMuPDF no disponible o falló: {exc}")

    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        pages = []
        for idx, page in enumerate(reader.pages, start=1):
            pages.append(f"\n\n<!-- page:{idx} -->\n\n{page.extract_text() or ''}")
        return "\n".join(pages), warnings
    except Exception as exc:
        warnings.append(f"pypdf no disponible o falló: {exc}")
        raise SystemExit(
            "No se pudo extraer texto del PDF. Instala una dependencia:\n"
            "  pip install pymupdf\n"
            "o:\n"
            "  pip install pypdf"
        )


def read_docx(path: Path) -> tuple[str, list[str]]:
    try:
        import docx  # type: ignore
    except Exception as exc:
        raise SystemExit("No se pudo leer DOCX. Instala: pip install python-docx") from exc

    document = docx.Document(str(path))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "").lower()
        if "heading 1" in style or "título 1" in style:
            blocks.append(f"# {text}")
        elif "heading 2" in style or "título 2" in style:
            blocks.append(f"## {text}")
        elif "heading 3" in style or "título 3" in style:
            blocks.append(f"### {text}")
        else:
            blocks.append(text)

    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [normalize_text(cell.text).replace("\n", " ") for cell in row.cells]
            rows.append(cells)
        if rows:
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            blocks.append("")
            blocks.append("| " + " | ".join(rows[0]) + " |")
            blocks.append("| " + " | ".join(["---"] * width) + " |")
            for row in rows[1:]:
                blocks.append("| " + " | ".join(row) + " |")
    return "\n\n".join(blocks), []


def read_html(path: Path) -> tuple[str, list[str]]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)</h1>", "\n\n", raw)
    raw = re.sub(r"(?i)<h1[^>]*>", "\n\n# ", raw)
    raw = re.sub(r"(?i)</h2>", "\n\n", raw)
    raw = re.sub(r"(?i)<h2[^>]*>", "\n\n## ", raw)
    raw = re.sub(r"(?i)</h3>", "\n\n", raw)
    raw = re.sub(r"(?i)<h3[^>]*>", "\n\n### ", raw)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</p>", "\n\n", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    return html.unescape(raw), []


def read_text_file(path: Path) -> tuple[str, list[str]]:
    return path.read_text(encoding="utf-8", errors="replace"), []


def convert(path: Path) -> Conversion:
    suffix = path.suffix.lower()
    warnings: list[str]
    if suffix == ".pdf":
        text, warnings = read_pdf(path)
        source_type = "pdf"
    elif suffix == ".docx":
        text, warnings = read_docx(path)
        source_type = "docx"
    elif suffix in {".html", ".htm"}:
        text, warnings = read_html(path)
        source_type = "html"
    elif suffix in {".md", ".markdown"}:
        text, warnings = read_text_file(path)
        source_type = "markdown"
    else:
        text, warnings = read_text_file(path)
        source_type = "text"

    markdown = normalize_text(text)
    return Conversion(
        source=str(path),
        source_type=source_type,
        extracted_at=datetime.now().isoformat(timespec="seconds"),
        chars=len(markdown),
        words=len(re.findall(r"\b\w+\b", markdown, flags=re.UNICODE)),
        warnings=warnings,
        markdown=markdown,
    )


def render_report(result: Conversion, include_text: bool = True) -> str:
    lines = [
        "# Documento convertido a Markdown",
        "",
        f"- Fuente: `{result.source}`",
        f"- Tipo: `{result.source_type}`",
        f"- Caracteres: {result.chars}",
        f"- Palabras aproximadas: {result.words}",
        f"- Extraído: {result.extracted_at}",
    ]
    if result.warnings:
        lines.append("- Advertencias:")
        lines.extend(f"  - {warning}" for warning in result.warnings)
    if include_text:
        lines.extend(["", "## Texto", "", result.markdown])
    return "\n".join(lines).rstrip() + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description="Convierte PDF, DOCX, HTML, TXT o MD a Markdown legible.")
    parser.add_argument("input", help="Archivo de entrada")
    parser.add_argument("--out", help="Ruta Markdown de salida")
    parser.add_argument("--json-out", help="Ruta JSON de salida")
    parser.add_argument("--no-text", action="store_true", help="No incluir el texto completo en el reporte Markdown")
    parser.add_argument('--overwrite', action='store_true', help='Reemplazar informes existentes, nunca entradas')
    args = parser.parse_args()
    try:
        validate_outputs([args.input], [args.out, args.json_out], args.overwrite)
    except ValueError as exc:
        parser.error(str(exc))

    path = Path(args.input)
    if not path.exists():
        raise SystemExit(f"No existe el archivo: {path}")

    result = convert(path)
    report = render_report(result, include_text=not args.no_text)

    if args.out:
        atomic_write(args.out, report, overwrite=args.overwrite)
    else:
        sys.stdout.write(report)

    if args.json_out:
        payload = asdict(result)
        atomic_write(args.json_out, json.dumps(payload, ensure_ascii=False, indent=2), overwrite=args.overwrite)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
