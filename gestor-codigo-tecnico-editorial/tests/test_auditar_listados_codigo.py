import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "auditar_listados_codigo.py"
DOCX_SCRIPT = ROOT / "scripts" / "generar_listado_docx.py"


class AuditarListadosCodigoTests(unittest.TestCase):
    @unittest.skipIf(os.name == "nt", "El wrapper POSIX se prueba en Linux/macOS")
    def test_posix_wrapper_runs_with_the_bundled_python(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "ejemplo.py"
            output = Path(directory) / "listado.docx"
            source.write_text("print('listado portable')\n", encoding="utf-8")
            completed = subprocess.run(
                ["sh", str(ROOT / "ejecutar.sh"), "generar-listado", str(source),
                 "--out", str(output), "--titulo", "Listado 1.4. Wrapper portable."],
                text=True,
                capture_output=True,
                check=False,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertTrue(output.is_file())

    @unittest.skipUnless(importlib.util.find_spec("docx"), "Requiere python-docx")
    def test_generates_a_native_docx_listing_without_libreoffice(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "validar_nombre.py"
            output = Path(directory) / "listado.docx"
            source.write_text(
                "def validar_nombre(nombre: str) -> str:\n"
                "    nombre_limpio = nombre.strip()\n"
                "    if not nombre_limpio:\n"
                "        raise ValueError('El nombre es obligatorio.')\n"
                "    return nombre_limpio\n",
                encoding="utf-8",
            )
            completed = subprocess.run(
                [sys.executable, str(DOCX_SCRIPT), str(source), "--out", str(output),
                 "--titulo", "Listado 1.1. Validación de un nombre."],
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertTrue(output.is_file())
            from docx import Document
            document = Document(output)
            self.assertEqual(document.paragraphs[0].text, "Listado 1.1. Validación de un nombre.")
            self.assertEqual(document.paragraphs[1].text, source.read_text(encoding="utf-8"))
            self.assertEqual(document.paragraphs[1].runs[0].font.name, "Consolas")

    def test_reports_untagged_uncaptioned_and_working_code_blocks(self):
        manuscript = """# Capítulo 2\n\n```\nprint('sin contexto editorial')\n```\n\n```python\n# TODO: completar ejemplo\nprint('borrador')\n```\n\n**Listado 2.1. Ejemplo correcto.**\n\n```bash\necho listo\n```\n"""
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "capitulo.md"
            report = Path(directory) / "reporte.json"
            source.write_text(manuscript, encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--json-out", str(report)],
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            findings = json.loads(report.read_text(encoding="utf-8"))
            by_check = {item["check"]: item for item in findings}
            self.assertEqual(by_check["bloques_sin_lenguaje"]["details"]["lines"], [3])
            self.assertEqual(by_check["bloques_sin_titulo"]["details"]["lines"], [3, 7])
            self.assertEqual(by_check["marcadores_de_trabajo"]["details"]["lines"], [8])


if __name__ == "__main__":
    unittest.main()
