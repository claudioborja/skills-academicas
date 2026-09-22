#!/usr/bin/env python3
"""Generar un listado DOCX nativo y breve sin depender de LibreOffice."""
from __future__ import annotations

import argparse
import io
from pathlib import Path
from archivos_seguros import atomic_output, validate_outputs

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
except ImportError as error:
    raise SystemExit(
        "Falta python-docx. Ejecuta ./ejecutar.sh generar-listado (Linux/macOS) o "
        ".\\ejecutar.ps1 generar-listado (Windows) para usar el intérprete empaquetado."
    ) from error


def build_document(code: str, title: str, language: str | None) -> Document:
    """Crear el documento con un bloque editable que conserva el código literal."""
    document = Document()
    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(6)
    title_run = title_paragraph.add_run(title)
    title_run.bold = True

    listing = document.add_paragraph()
    listing.paragraph_format.space_before = Pt(0)
    listing.paragraph_format.space_after = Pt(6)
    listing.paragraph_format.line_spacing = 1
    listing.paragraph_format.keep_together = True
    code_run = listing.add_run(code)
    code_run.font.name = "Consolas"
    code_run.font.size = Pt(9.5)

    if language:
        note = document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        note.paragraph_format.space_before = Pt(0)
        note.paragraph_format.space_after = Pt(0)
        note.add_run(f"Lenguaje: {language}.").italic = True
    return document


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Archivo de código UTF-8 que se publicará literalmente.")
    parser.add_argument("--out", required=True, type=Path, help="DOCX de salida; no se sobrescribe por defecto.")
    parser.add_argument("--titulo", required=True, help="Rótulo editorial, por ejemplo: Listado 1.1. Validación.")
    parser.add_argument("--lenguaje", help="Lenguaje que se declara debajo del listado.")
    parser.add_argument("--overwrite", action="store_true", help="Autorizar reemplazo del DOCX de salida.")
    args = parser.parse_args()
    try:
        source = args.input.resolve()
        if not source.is_file():
            raise ValueError(f"No existe el archivo de código: {args.input}")
        output = args.out.resolve()
        validate_outputs([source], [output], overwrite=args.overwrite)
        code = source.read_text(encoding="utf-8")
        document = build_document(code, args.titulo, args.lenguaje)
        payload = io.BytesIO()
        document.save(payload)
        with atomic_output(output, overwrite=args.overwrite) as target:
            target.write(payload.getvalue())
    except (OSError, UnicodeError, ValueError) as error:
        parser.error(str(error))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
